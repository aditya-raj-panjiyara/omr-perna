"""
Build a student result sheet from OMR answers in Excel and an answer key.

  .venv/bin/python make_results.py                      (finds the answers Excel and answer key itself)
  .venv/bin/python make_results.py answers.xlsx answer_key.csv
  .venv/bin/python make_results.py answers.xlsx answer_key.xlsx -o results.xlsx --correct 2 --wrong -0.5

answers  : .xlsx/.xls/.csv, one row per student. Question columns are named q1..q100, Q1..Q100 or 1..100;
           every other column (Roll, Name, ...) is copied to the output as-is.
           Each answer is A/B/C/D. X (or two letters such as AB) = multiple marked -> incorrect.
           Blank, O, - = not attempted.
key      : .csv/.xlsx with lines like "q1,B" (or "1,B"), or a single row with q1..q100 headers.

Output (default: <answers>_results.xlsx) has three sheets:
  Results   - per student: Correct / Incorrect / Not attempted / Marks for every subject,
              Paper 1, Paper 2 and Total, plus Rank
  Answers   - every student's answers, green = correct, red = incorrect, grey = not attempted
  Answer Key
"""
import argparse
import os
import glob
import re
import shutil
import sys

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

N_QUESTIONS = 100
SUBJECTS = [  # (name, first question, last question)
    ("Reasoning (MAT)", 1, 50),
    ("Social Science", 51, 75),
    ("Chemistry", 76, 82),
    ("Biology", 83, 94),
    ("Physics", 95, 100),
]
PAPERS = [("Paper 1 (MAT)", 1, 50), ("Paper 2", 51, 100), ("Total", 1, 100)]
STATS = ["Correct", "Incorrect", "Not Attempted", "Marks"]
OMR_SUMMARY_COLS = ["Answered", "Blank", "Correct", "Wrong", "Score"]  # omr.py's own totals, recomputed here
BLANK_VALUES = {"", "O", "-", "NAN", "NONE"}
Q_HEADER = re.compile(r"^\s*q?\s*\.?\s*(\d{1,3})\s*$", re.I)


def question_number(header):
    if isinstance(header, (int, float)) and not pd.isna(header) and float(header).is_integer():
        n = int(header)
    else:
        m = Q_HEADER.match(str(header))
        n = int(m.group(1)) if m else None
    return n if n and 1 <= n <= N_QUESTIONS else None


def read_table(path, header="infer"):
    if path.lower().endswith((".xlsx", ".xlsm", ".xls")):
        return pd.read_excel(path, header=0 if header == "infer" else header, dtype=str)
    return pd.read_csv(path, header=header, dtype=str)


def answers_table(path):
    """CSV, or the first Excel sheet whose header row has all 100 question columns (e.g. 'OMR_results')."""
    if not path.lower().endswith((".xlsx", ".xlsm", ".xls")):
        return read_table(path)
    sheets = pd.read_excel(path, sheet_name=None, dtype=str)
    for df in sheets.values():
        if len({question_number(c) for c in df.columns} - {None}) == N_QUESTIONS:
            return df
    return next(iter(sheets.values()))


def load_answers(path):
    df = answers_table(path).fillna("")
    qcols = {}
    for c in df.columns:
        n = question_number(c)
        if n and n not in qcols:
            qcols[n] = c
    missing = [n for n in range(1, N_QUESTIONS + 1) if n not in qcols]
    if missing:
        sys.exit(f"Answer sheet is missing question columns: {missing[:10]}{' ...' if len(missing) > 10 else ''}\n"
                 f"Columns found: {list(df.columns)[:15]} ...")
    info_cols = [c for c in df.columns if c not in qcols.values()]
    answers = df[[qcols[n] for n in range(1, N_QUESTIONS + 1)]].copy()
    answers.columns = range(1, N_QUESTIONS + 1)
    answers = answers.apply(lambda col: col.str.strip().str.upper())
    if "MultiMarked" in df.columns:  # output of omr.py: double-marked answers were cleared, put them back as X
        for i, qs in df["MultiMarked"].items():
            for q in str(qs).split():
                n = question_number(q)
                if n:
                    answers.at[i, n] = "X"
    info_cols = [c for c in info_cols if c not in OMR_SUMMARY_COLS + ["School Code"]]
    df = df[info_cols].copy()
    if "Sheet" in df.columns:  # each PDF is one school: "(29) 8TH_p05" -> school 29
        code = pd.to_numeric(df["Sheet"].str.extract(r"^\s*\((\d+)\)", expand=False), errors="coerce")
        df.insert(df.columns.get_loc("Sheet") + 1, "School Code", code.astype("Int64"))
    df = df.loc[:, [c for c in df.columns if not str(c).startswith("Unnamed")] or info_cols]
    return df, answers


def load_key(path):
    raw = read_table(path, header=None).fillna("")
    key = {}
    for row in raw.itertuples(index=False):  # long form: "q1,B" per line
        if len(row) >= 2:
            n, a = question_number(row[0]), str(row[1]).strip().upper()
            if n and a in "ABCD" and len(a) == 1:
                key[n] = a
    if len(key) < N_QUESTIONS and len(raw) >= 2:  # wide form: header row of questions, answers below
        for header, a in zip(raw.iloc[0], raw.iloc[1]):
            n, a = question_number(header), str(a).strip().upper()
            if n and len(a) == 1 and a in "ABCD":
                key[n] = a
    missing = [n for n in range(1, N_QUESTIONS + 1) if n not in key]
    if missing:
        sys.exit(f"Answer key has no A/B/C/D answer for questions: {missing}")
    return key


def verdict(answer, correct):
    if answer in BLANK_VALUES:
        return "N"
    return "C" if answer == correct else "I"  # X, AB, etc. never equal a single key letter


def build(info, answers, key, marks):
    v = pd.DataFrame({n: [verdict(a, key[n]) for a in answers[n]] for n in answers.columns}, index=answers.index)
    out = info.copy()
    for name, lo, hi in SUBJECTS + PAPERS:
        part = v.loc[:, lo:hi]
        c, i, n = (part == "C").sum(axis=1), (part == "I").sum(axis=1), (part == "N").sum(axis=1)
        out[(name, "Correct")] = c
        out[(name, "Incorrect")] = i
        out[(name, "Not Attempted")] = n
        out[(name, "Marks")] = (c * marks["correct"] + i * marks["wrong"] + n * marks["blank"]).round(2)
    total = out[("Total", "Marks")]
    out[("Total", "Rank")] = total.rank(ascending=False, method="min").astype(int)
    if "School Code" in out.columns:
        out[("Total", "School Rank")] = total.groupby(out["School Code"]).rank(ascending=False, method="min").astype("Int64")
        order = out.sort_values(["School Code", ("Total", "School Rank"), "Sheet"]).index
    else:
        order = out.sort_values(("Total", "Rank")).index
    return out.loc[order], v.loc[order]


def total_extras(results):
    return [c for c in ("Rank", "School Rank") if ("Total", c) in results.columns]


def write_excel(path, info, results, answers, verdicts, key, marks, update=False):
    info, answers = info.loc[results.index], answers.loc[results.index]
    info_cols = list(info.columns)
    groups = [name for name, _, _ in SUBJECTS + PAPERS]
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    sub_fill = PatternFill("solid", fgColor="DDEBF7")
    total_fill = PatternFill("solid", fgColor="FFF2CC")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    mode = dict(mode="a", if_sheet_exists="replace") if update else dict(mode="w")
    with pd.ExcelWriter(path, engine="openpyxl", **mode) as xw:
        # --- Results: two header rows (subject, then stat) ---
        flat = results[info_cols].copy()
        for g in groups:
            for s in STATS + (total_extras(results) if g == "Total" else []):
                flat[f"{g} | {s}"] = results[(g, s)]
        flat.to_excel(xw, sheet_name="Results", index=False, startrow=1)
        ws = xw.sheets["Results"]
        col = 1
        for c in info_cols:
            ws.cell(1, col, str(c))
            ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
            col += 1
        for g in groups:
            stats = STATS + (total_extras(results) if g == "Total" else [])
            ws.cell(1, col, g)
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(stats) - 1)
            for k, s in enumerate(stats):
                ws.cell(2, col + k, s)
                if g == "Total":
                    for r in range(3, ws.max_row + 1):
                        ws.cell(r, col + k).fill = total_fill
            col += len(stats)
        for r in (1, 2):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(r, c)
                cell.alignment, cell.border = center, box
                cell.font = Font(bold=True, color="FFFFFF" if r == 1 else "000000")
                cell.fill = head_fill if r == 1 else sub_fill
        for row in ws.iter_rows(min_row=3):
            for cell in row:
                cell.border = box
                if cell.column > len(info_cols):
                    cell.alignment = Alignment(horizontal="center")
        for c in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(c)].width = 16 if c <= len(info_cols) else 10
        ws.row_dimensions[2].height = 30
        ws.freeze_panes = ws.cell(3, len(info_cols) + 1)

        # --- Answers: what each student marked, coloured by verdict ---
        sheet = info.copy()
        for n in range(1, N_QUESTIONS + 1):
            sheet[f"Q{n}"] = answers[n]
        sheet.to_excel(xw, sheet_name="Answers", index=False, startrow=1)
        ws = xw.sheets["Answers"]
        ws.cell(1, 1, "Key →").font = Font(bold=True)
        fills = {"C": PatternFill("solid", fgColor="C6EFCE"), "I": PatternFill("solid", fgColor="FFC7CE"),
                 "N": PatternFill("solid", fgColor="EDEDED")}
        for n in range(1, N_QUESTIONS + 1):
            c = len(info_cols) + n
            ws.cell(1, c, key[n]).font = Font(bold=True)
            ws.cell(1, c).alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(c)].width = 5
            for r, v in enumerate(verdicts[n], start=3):
                ws.cell(r, c).fill = fills[v]
                ws.cell(r, c).alignment = Alignment(horizontal="center")
        ws.freeze_panes = ws.cell(3, len(info_cols) + 1)

        # --- Answer key + marking used ---
        ks = pd.DataFrame({"Question": range(1, N_QUESTIONS + 1), "Answer": [key[n] for n in range(1, N_QUESTIONS + 1)],
                           "Subject": [next(s for s, lo, hi in SUBJECTS if lo <= n <= hi) for n in range(1, N_QUESTIONS + 1)]})
        ks.to_excel(xw, sheet_name="Answer Key", index=False)
        ws = xw.sheets["Answer Key"]
        ws.cell(1, 5, "Marking").font = Font(bold=True)
        for r, (k, label) in enumerate([("correct", "Correct"), ("wrong", "Incorrect"), ("blank", "Not attempted")], 2):
            ws.cell(r, 5, label)
            ws.cell(r, 6, marks[k])
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["E"].width = 14


KIT = os.path.dirname(os.path.abspath(__file__))
SEARCH_DIRS = [KIT, "~/Desktop", "~/Downloads", "~/Documents"]
SKIP_DIRS = {".venv", "OMRChecker", "_omrchecker", "_work", "checked_sheets", "Library", "node_modules"}


def candidate_files(exts, max_depth=3):
    for base in SEARCH_DIRS:
        base = os.path.expanduser(base)
        for root, dirs, files in os.walk(base):
            dirs[:] = [d for d in dirs if d not in SKIP_DIRS and not d.startswith(".")]
            if root[len(base):].count(os.sep) >= max_depth:
                dirs[:] = []
            for f in files:
                if f.lower().endswith(exts) and not f.startswith(("~$", ".")):
                    yield os.path.join(root, f)


def has_question_columns(path):
    try:
        cols = read_table(path).columns if path.lower().endswith(".csv") else pd.read_excel(path, nrows=0).columns
    except Exception:
        return False
    return len({question_number(c) for c in cols} - {None}) == N_QUESTIONS


def find_answers():
    """Newest Excel/CSV (not one of our own outputs) that has columns for all 100 questions."""
    found = sorted({os.path.realpath(f) for f in candidate_files((".xlsx", ".xls", ".csv"))
                    if not f.endswith(("_results.xlsx", "Student_Results.xlsx")) and has_question_columns(f)}, key=os.path.getmtime, reverse=True)
    if not found:
        sys.exit("Could not find a students' answers file (an Excel with columns q1..q100).\n"
                 "Put it in the omr_kit folder, Desktop, Downloads or Documents, or give its path:\n"
                 "  .venv/bin/python make_results.py \"/path/to/answers.xlsx\" answer_key_class8.csv")
    if len(found) > 1:
        print("Several answer files found, using the newest. Others:\n  " + "\n  ".join(found[1:6]))
    return found[0]


def find_key():
    keys = sorted(f for f in glob.glob(os.path.join(KIT, "answer_key*.csv")) if "TEMPLATE" not in f)
    if len(keys) != 1:
        sys.exit("Give the answer key file after the answers file. Answer keys in the kit: " + (", ".join(map(os.path.basename, keys)) or "none"))
    return keys[0]


def main():
    ap = argparse.ArgumentParser(description="Generate student results (subject, paper and total) from OMR answers.")
    ap.add_argument("answers", nargs="?", help="Excel/CSV with one row per student and columns q1..q100 (found automatically if left out)")
    ap.add_argument("key", nargs="?", help="answer key: CSV/Excel with lines like q1,B (found automatically if left out)")
    ap.add_argument("-o", "--out", help="output .xlsx (default: <answers>_results.xlsx)")
    ap.add_argument("--update", metavar="XLSX", help="write Results/Answers/Answer Key into this existing workbook "
                    "(other sheets are kept; a .backup copy is made first)")
    ap.add_argument("--correct", type=float, default=1, help="marks for a correct answer (default 1)")
    ap.add_argument("--wrong", type=float, default=0, help="marks for an incorrect answer, e.g. -0.25 (default 0)")
    ap.add_argument("--blank", type=float, default=0, help="marks for not attempted (default 0)")
    args = ap.parse_args()
    args.answers = args.answers or find_answers()
    args.key = args.key or find_key()
    print(f"Student answers: {args.answers}\nAnswer sheet:    {args.key}")

    for p in (args.answers, args.key):
        if not os.path.isfile(p):
            sys.exit(f"File not found: {p}\nTip: drag the file from Finder into Terminal to paste its full path.")
    marks = {"correct": args.correct, "wrong": args.wrong, "blank": args.blank}
    info, answers = load_answers(args.answers)
    key = load_key(args.key)
    results, verdicts = build(info, answers, key, marks)
    stem = os.path.splitext(args.answers)[0]
    out = args.out or (os.path.join(os.path.dirname(stem), "Student_Results.xlsx") if os.path.basename(stem) == "OMR_results"
                       else stem + "_results.xlsx")
    if args.update:
        out = args.update
        backup = os.path.splitext(out)[0] + ".backup.xlsx"
        shutil.copy2(out, backup)
        print(f"Backup of the original: {backup}")
    write_excel(out, info, results, answers, verdicts, key, marks, update=bool(args.update))

    odd = sorted({a for col in answers.columns for a in answers[col]} - BLANK_VALUES - set("ABCD") - {"X"})
    if odd:
        print(f"Note: unusual answer values counted as incorrect: {odd[:20]}")
    t = results[("Total", "Marks")]
    print(f"{len(results)} students | total marks: avg {t.mean():.1f}, max {t.max():g}, min {t.min():g}")
    if "School Code" in results.columns:
        print(f"{results['School Code'].nunique()} schools: " + ", ".join(
            f"{c} ({n})" for c, n in results["School Code"].value_counts().sort_index().items()))
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
